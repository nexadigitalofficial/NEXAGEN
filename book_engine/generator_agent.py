import os
import sys
import time
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

OUTPUT_DIR = r"C:\Users\USER\Desktop\kitap"
os.makedirs(OUTPUT_DIR, exist_ok=True)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_styled_document():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Style definitions
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Georgia'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal_style.paragraph_format.line_spacing = 1.35
    normal_style.paragraph_format.space_after = Pt(6)

    return doc

def add_title_page(doc, title, subtitle, author="NEXAGEN OMEGA Çok-Ajanlı Bilim Ordusu"):
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(80)
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(f"{title}\n")
    r_title.font.name = 'Georgia'
    r_title.font.size = Pt(26)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x00, 0x2B, 0x5B) # Deep Navy
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run(f"{subtitle}\n\n")
    r_sub.font.name = 'Georgia'
    r_sub.font.size = Pt(14)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0x80, 0x00, 0x20) # Burgundy
    
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_div = p_div.add_run("―" * 35)
    r_div.font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA)
    
    p_author = doc.add_paragraph()
    p_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_author.paragraph_format.space_before = Pt(120)
    r_author = p_author.add_run(f"{author}\n")
    r_author.font.size = Pt(12)
    r_author.font.bold = True
    r_author.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    r_subauthor = p_author.add_run("Kürsü Yayını: Moleküler Nörogenetik, Biyofizik ve Sentetik Kognisyon Serisi\n2026 Resmî Akademik Referans Eseri")
    r_subauthor.font.size = Pt(10)
    r_subauthor.font.italic = True
    r_subauthor.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    
    doc.add_page_break()

def add_callout_box(doc, title, text_content, border_color="002B5B", bg_color="F0F4F8"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    r_title = p.add_run(f"{title}\n")
    r_title.font.bold = True
    r_title.font.size = Pt(11)
    r_title.font.color.rgb = RGBColor(0x00, 0x2B, 0x5B)
    
    r_body = p.add_run(text_content)
    r_body.font.size = Pt(10)
    r_body.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_table_data(doc, headers, rows_data):
    tbl = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = True
    
    # Header Row
    hdr_cells = tbl.rows[0].cells
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text
        set_cell_background(hdr_cells[i], "002B5B")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            
    # Data Rows
    for r_idx, row in enumerate(rows_data):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_hex = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_hex)
            set_cell_margins(row_cells[c_idx], top=100, bottom=100, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            for run in p.runs:
                run.font.size = Pt(9.0)
                run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

print("[NEXAGEN OMEGA] Generator Agent Module Initialized.")
